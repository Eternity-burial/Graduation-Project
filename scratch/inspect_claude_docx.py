import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc_path = r"D:\tj\Graduation Project\开题报告\开题报告_claude版.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check formulas
omaths = doc._element.xpath('.//m:oMath')
omathparas = doc._element.xpath('.//m:oMathPara')
print(f"Total oMath elements: {len(omaths)}")
print(f"Total oMathPara (display formulas): {len(omathparas)}")

# Print paragraph summary
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if p._element.xpath('.//m:oMathPara'):
        print(f"P{i:02d} [DISPLAY FORMULA]: {t}")
    elif p.text.startswith("一、") or p.text.startswith("二、") or p.text.startswith("三、") or p.text.startswith("四、"):
        print(f"P{i:02d} [H1]: {t}")
    elif p.text.startswith("1．") or p.text.startswith("2．") or p.text.startswith("3．"):
        print(f"P{i:02d} [H2]: {t}")
    elif p.text.startswith("2.1") or p.text.startswith("2.2") or p.text.startswith("2.3") or p.text.startswith("2.4"):
        print(f"P{i:02d} [H3]: {t}")
