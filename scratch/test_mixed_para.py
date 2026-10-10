import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn
import lxml.etree as etree
import latex2mathml.converter
import win32com.client as win32

xsl_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = etree.parse(xsl_path)
transform = etree.XSLT(xslt)

def latex_to_omml(latex, is_display=False):
    mathml = latex2mathml.converter.convert(latex)
    dom = etree.fromstring(mathml)
    new_dom = transform(dom)
    omml_str = etree.tostring(new_dom).decode('utf-8')
    if is_display:
        return parse_xml(f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{omml_str}</m:oMathPara>')
    else:
        return parse_xml(omml_str)

doc = docx.Document()
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
sp = OxmlElement('w:spacing')
sp.set(qn('w:line'), '360')
sp.set(qn('w:lineRule'), 'auto')
pPr.append(sp)

# Add run
r1 = p.add_run('其中，')
r1.font.name = 'Times New Roman'
r1._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/>'))

# Add inline math
p._p.append(latex_to_omml(r'M = M_{RB} + M_A \in \mathbb{R}^{6\times6}'))

# Add run
r2 = p.add_run(' 为含附加质量 ')
r2.font.name = 'Times New Roman'
r2._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/>'))

p._p.append(latex_to_omml(r'M_A'))

r3 = p.add_run(' 的系统惯性矩阵，严卫生等')
r3.font.name = 'Times New Roman'
r3._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/>'))

# Superscript citation
r4 = p.add_run('[6]')
r4.font.name = 'Times New Roman'
r4.font.superscript = True

test_path = os.path.abspath('scratch/test_mixed_para.docx')
doc.save(test_path)

word = win32.gencache.EnsureDispatch('Word.Application')
word.Visible = False
try:
    com_doc = word.Documents.Open(test_path, ReadOnly=True)
    print('Successfully opened in Word COM! Paragraphs count:', com_doc.Paragraphs.Count)
    com_doc.Close(False)
finally:
    word.Quit()
print('Word COM test PASSED!')
