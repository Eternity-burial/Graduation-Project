import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc_path = r"开题报告\开题报告_claude版.docx"
doc = docx.Document(doc_path)

print("=== TABLE 0 (COVER) ===")
for r_idx, row in enumerate(doc.tables[0].rows):
    row_text = [cell.text.strip() for cell in row.cells]
    print(f"Row {r_idx}: {row_text}")

print("\n=== TABLE 1 (DATE) ===")
date_text = [cell.text.strip() for cell in doc.tables[1].rows[0].cells]
print(f"Date row: {date_text}")

print("\n=== SAMPLE SECTION 1.1 ===")
print("P17:", doc.paragraphs[17].text)
print("P18 (Body 1):", doc.paragraphs[18].text[:80])
print("P19 (Body 2):", doc.paragraphs[19].text[:80])
print("P20 (Body 3):", doc.paragraphs[20].text[:80])

print("\n=== SAMPLE SECTION 1.2 ===")
print("P21:", doc.paragraphs[21].text)
print("P22 (Lead):", doc.paragraphs[22].text[:80])
print("P23 (H3 2.1):", doc.paragraphs[23].text)
print("P24 (Para 1):", doc.paragraphs[24].text[:80])
print("P25 (Lead-in):", doc.paragraphs[25].text[:80])
print("P26 (Formula 1 xml):", [e.tag.split('}')[-1] for e in doc.paragraphs[26]._p])
print("P27 (Para with inline math):", doc.paragraphs[27].text[:80])

print("\n=== SAMPLE SECTION 2 ===")
print("P72:", doc.paragraphs[72].text)
print("P73:", doc.paragraphs[73].text)
print("P74:", doc.paragraphs[74].text)
print("P75:", doc.paragraphs[75].text)
print("P76:", doc.paragraphs[76].text)
print("P77:", doc.paragraphs[77].text)
print("P78:", doc.paragraphs[78].text)

print("\n=== SAMPLE SECTION 3 (REFS) ===")
print("P79:", doc.paragraphs[79].text)
print("P80 (Ref 1):", doc.paragraphs[80].text[:80])
print("P81 (Ref 2):", doc.paragraphs[81].text[:80])
print("P116 (Ref 37):", doc.paragraphs[116].text[:80])

print("\n=== SAMPLE SECTION 4 ===")
print("P117:", doc.paragraphs[117].text)
print("Table 2 (Audits):", len(doc.tables[2].rows), "rows")
