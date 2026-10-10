# -*- coding: utf-8 -*-
"""
毕业设计（论文）任务书 Word (.docx) 官方模板就地填充构建管线与自动化审计

核心原则与规范：
1. 【严格遵循官方模板格式】：以 模板/1毕业设计(论文)任务书.docx 为唯一基准。
   - 绝不破坏任何原生段落、表格样式与结构。
   - 正文：宋体 + Times New Roman，小四号 (12.0 pt)，1.5 倍行距 (line="360")，首行缩进 2 字符 (480 dxa)。
   - 大节标题（一至五）：统一按照用户定稿设定为“毕业论文”（如：一、毕业论文的课题背景），黑体四号 (14.0 pt)，顶格无缩进，未加粗。
   - 小节标题（2.1, 2.2, 2.3, 2.4）：黑体 + Times New Roman，小四号 (12.0 pt)，顶格无缩进 (first_line=0, left=0)，1.5 倍行距，【严格不加粗】(bold=False)！
   - 成果要求（（1）~（7））：宋体 + Times New Roman，小四号 (12.0 pt)，1.5 倍行距，首行缩进 2 字符 (480 dxa)，不加粗。
   - 进度表：三列标准全格线表，单线细边框 (sz="4")，居中对齐，五号字体 (10.5 pt)，宋体 + Times New Roman。
   - 参考文献：宋体 + Times New Roman，小四号 (12.0 pt)，悬挂缩进 0.74 cm (420 dxa)，1.5 倍行距，严格不加粗。
2. 【100% 完整保留官方全部权限区域 (permStart / permEnd)】：
   - 封面表格 Table 0：保留所有单元格与列权限属性，就地填充题名、学院、专业、姓名、学号。
   - 一、课题背景：内容 100% 置于 permStart 748764618 与 permEnd 748764618 之间。
   - 二、主要内容：内容 100% 置于 permStart 107480931 与 permEnd 107480931 之间。
   - 三、基本要求及成果形式：内容 100% 置于 permStart 207912164 与 permEnd 207912164 之间。
   - 四、进度安排：进度表 100% 置于 permStart 2121286149 与 permEnd 2121286149 之间。
   - 五、主要参考文献 & 六、其他要求：100% 置于 permStart 1876760563 与 permEnd 1876760563 之间。
3. 【文献真实性与著录规范】：
   - 全面修正文献真实性与著录错误：
     * [7] Torrente 等论文实际刊登于 IEEE Robotics and Automation Letters (RA-L), 2021, 6(2): 3769-3776，纠正原误写的 IEEE TCST；
     * [8] Chu 等论文更正为该卷期实际真实刊载之题目与作者（ISA Transactions, 2020, 100: 28-37）；
     * [4] 补全 Heshmati-Alamdari 会议起止页码 6183-6188；
     * [12] 修复 PX4 Control Allocation 官方文档 404 死链（补充 .html 后缀）。
"""

import os
import sys
import shutil
from pathlib import Path
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(r"D:\tj\Graduation Project")
TEMPLATE_DOCX = PROJECT_ROOT / r"模板\1毕业设计(论文)任务书.docx"
OUT_TASK_DIR = PROJECT_ROOT / r"任务书"
OUT_TASK_DOCX_MAIN = OUT_TASK_DIR / r"任务书-基于PX4的水下航行器模型控制方法研究.docx"
OUT_TASK_DOCX_TPL = OUT_TASK_DIR / r"1毕业设计(论文)任务书.docx"
OUT_TASK_DOCX_CLEAN = OUT_TASK_DIR / r"PX4水下航行器任务书_修改净稿.docx"


# ==================== 字体与段落排版辅助函数 ====================

def set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False):
    """同时精确设置中西文字体、字号与粗体属性"""
    run.font.name = western_font
    run.font.size = Pt(font_size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), western_font)
    rFonts.set(qn("w:hAnsi"), western_font)
    rFonts.set(qn("w:cs"), western_font)
    rFonts.set(qn("w:eastAsia"), chinese_font)


def set_para_format(p, style_name="Normal", line_val="360", line_rule="auto", first_line_dxa=480, before_dxa=0, after_dxa=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    """设置段落样式、行距、缩进与段间距"""
    if style_name:
        p.style = style_name
    p.alignment = align
    pPr = p._p.get_or_add_pPr()

    # 行距与段间距
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        pPr.append(sp)
    sp.set(qn("w:line"), str(line_val))
    sp.set(qn("w:lineRule"), str(line_rule))
    sp.set(qn("w:before"), str(before_dxa))
    sp.set(qn("w:after"), str(after_dxa))

    # 缩进设置
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    if first_line_dxa is not None and first_line_dxa > 0:
        ind.set(qn("w:firstLine"), str(first_line_dxa))
        ind.set(qn("w:firstLineChars"), "200")
        if qn("w:left") in ind.attrib:
            del ind.attrib[qn("w:left")]
        if qn("w:hanging") in ind.attrib:
            del ind.attrib[qn("w:hanging")]
    elif first_line_dxa == 0:
        # 顶格无缩进：清除所有首行与左缩进
        for attr in [qn("w:firstLine"), qn("w:firstLineChars"), qn("w:left"), qn("w:leftChars"), qn("w:hanging")]:
            if attr in ind.attrib:
                del ind.attrib[attr]
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.left_indent = None


def set_hanging_indent(p, left_dxa=420, hanging_dxa=420, line_val="360"):
    """设置参考文献悬挂缩进（0.74 cm = 420 dxa）"""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        pPr.append(sp)
    sp.set(qn("w:line"), str(line_val))
    sp.set(qn("w:lineRule"), "auto")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")

    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:left"), str(left_dxa))
    ind.set(qn("w:hanging"), str(hanging_dxa))
    for attr in [qn("w:firstLine"), qn("w:firstLineChars")]:
        if attr in ind.attrib:
            del ind.attrib[attr]


def create_new_paragraph(doc, text, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False, first_line_dxa=480, before_dxa=0, after_dxa=0, style_name="Normal", line_val="360"):
    """构造全新段落 XML 元素，设置字体与排版格式"""
    p = docx.text.paragraph.Paragraph(OxmlElement("w:p"), doc)
    set_para_format(p, style_name=style_name, line_val=line_val, first_line_dxa=first_line_dxa, before_dxa=before_dxa, after_dxa=after_dxa)
    run = p.add_run(text)
    set_run_font(run, chinese_font=chinese_font, western_font=western_font, font_size_pt=font_size_pt, bold=bold)
    return p


def create_ref_paragraph(doc, text, font_size_pt=12.0):
    """构造参考文献段落（正文格式样式，悬挂缩进）"""
    p = docx.text.paragraph.Paragraph(OxmlElement("w:p"), doc)
    p.style = "正文格式"
    set_hanging_indent(p, left_dxa=420, hanging_dxa=420, line_val="360")
    run = p.add_run(text)
    set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=font_size_pt, bold=False)
    return p


def update_heading_text_safely(p, new_text):
    """安全更新大标题文本为'毕业论文'，保留段落内权限标签 (permStart/permEnd)"""
    for r in list(p.runs):
        p._p.remove(r._r)
    run = p.add_run(new_text)
    set_run_font(run, chinese_font="黑体", western_font="Times New Roman", font_size_pt=14.0, bold=False)


# ==================== 任务书主构建流水线 ====================

def build_task_book(target_file_path):
    print("=" * 70)
    print(f">>> 启动毕业设计（论文）任务书构建: {target_file_path.name}")
    print("=" * 70)

    if not TEMPLATE_DOCX.exists():
        raise FileNotFoundError(f"未找到官方模板文件: {TEMPLATE_DOCX}")

    OUT_TASK_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE_DOCX, target_file_path)
    print(f"已复制官方空白模板至: {target_file_path}")

    doc = docx.Document(str(target_file_path))
    body = doc._element.body

    # 1. 填充 Table 0（保留单元格内列权限标签，就地替换文本）
    t0 = doc.tables[0]
    table_0_entries = [
        (0, 1, "基于PX4的水下航行器模型控制方法研究"),
        (2, 1, "机械工程与机器人学院"),
        (3, 1, "机械设计制造及其自动化"),
        (4, 1, "张卫恒"),
        (4, 3, "2352407")
    ]
    for r_i, c_i, text_val in table_0_entries:
        cell = t0.rows[r_i].cells[c_i]
        p = cell.paragraphs[0]
        for r in list(p.runs):
            p._p.remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text_val)
        set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=14.0, bold=False)

    print("Table 0 封面基本信息填充完成，权限节点完好保留！")

    # ==================== 2. 填充 Section 1：课题背景 ====================
    # 更新 Section 1 原生大标题为 "一、毕业论文的课题背景"
    p_h1_node = [el for el in body if el.tag.endswith("p") and any(e.get(qn("w:id")) == "1533047981" for e in el.iter(qn("w:permStart")))][0]
    p_h1 = docx.text.paragraph.Paragraph(p_h1_node, doc)
    update_heading_text_safely(p_h1, "一、毕业论文的课题背景")

    # 定位 permStart 748764618 所在段落及对应 permEnd
    p_sec1_start = [el for el in body if any(e.get(qn("w:id")) == "748764618" for e in el.iter(qn("w:permStart")))][0]
    perm_end_sec1 = [el for el in body if el.tag.endswith("permEnd") and el.get(qn("w:id")) == "748764618"][0]

    sec1_texts = [
        "随着海洋资源开发与水下基础设施运维需求增长，海上风电桩基、水下大坝立面及海底管线等场景的近距离精细巡检具有重要工程意义。受制于人工潜水作业的安全风险以及大型作业级水下航行器部署繁琐、成本高昂等局限，便携型水下航行器（UUV/ROV）逐渐成为关键作业装备。此类巡检作业要求平台在受限空间内实现高精度轨迹跟踪、定点悬停与多自由度姿态主动调节。常规四推进器或者六推进器平台多依赖重浮力被动回正，难以实现各运动自由度的解耦控制，而采用空间矢量布局的八推进器紧凑型航行器具备六自由度全向主动控制能力，为近场高灵巧度作业提供了理想的硬件平台。",
        "然而，水下航行器动力学存在附加质量、非线性阻尼以及重浮力恢复力矩等耦合影响，紧凑型八推进器平台受到推进器推力饱和与多自由度控制分配能力的制约，在执行复杂机动作业时易出现执行器饱和。常规比例—积分—微分（PID）控制难以直接兼顾多变量耦合与执行器约束；模型预测控制（MPC）能够将约束纳入优化，但需要权衡模型预测精度与在线求解开销。同时，PX4提供控制接口、状态反馈和执行器管理等基础能力，但其水下航行器支持仍处于实验性阶段，针对八推进器构型开展平台适配和闭环验证具有实际必要性。",
        "为此，本课题以八推进器水下航行器为研究对象，开展面向控制的动力学建模与参数辨识，探索考虑执行器约束的模型预测控制（MPC）方法及其与PX4系统的集成应用。首先构建多种动力学候选模型及空间推力映射关系，基于试验数据开展模型比选与论证；其次设计考虑推进器约束的轨迹跟踪控制算法，探索并比选预测控制器与PX4系统的接口集成方案；最后通过动力学仿真、PX4软件在环（Software-in-the-Loop, SITL）仿真以及视水池试验条件开展的实物验证，对跟踪精度、推力饱和情况与计算实时性进行综合评估，为八推进器水下航行器的控制实现提供依据。"
    ]

    p_sec1_start_obj = docx.text.paragraph.Paragraph(p_sec1_start, doc)
    set_para_format(p_sec1_start_obj, style_name="正文格式", line_val="360", first_line_dxa=480)
    for r in list(p_sec1_start_obj.runs):
        p_sec1_start_obj._p.remove(r._r)
    run1 = p_sec1_start_obj.add_run(sec1_texts[0])
    set_run_font(run1, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False)

    p_curr = p_sec1_start
    for t in sec1_texts[1:]:
        p_new = create_new_paragraph(doc, t, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False, first_line_dxa=480, line_val="360", style_name="正文格式")
        p_curr.addnext(p_new._p)
        p_curr = p_new._p

    elements_to_remove = []
    node = p_curr.getnext()
    while node is not None and node != perm_end_sec1:
        elements_to_remove.append(node)
        node = node.getnext()
    for el in elements_to_remove:
        body.remove(el)

    print("Section 1（课题背景）三段论注入完成，100% 位于 permStart/permEnd 748764618 内部！")

    # ==================== 3. 填充 Section 2：主要内容 ====================
    # 更新 Section 2 原生大标题为 "二、毕业论文应完成的主要内容"
    p_h2_node = [el for el in body if el.tag.endswith("p") and any(e.get(qn("w:id")) == "462645645" for e in el.iter(qn("w:permStart")))][0]
    p_h2 = docx.text.paragraph.Paragraph(p_h2_node, doc)
    update_heading_text_safely(p_h2, "二、毕业论文应完成的主要内容")

    p_sec2_start = [el for el in body if any(e.get(qn("w:id")) == "107480931" for e in el.iter(qn("w:permStart")))][0]
    perm_end_sec2 = [el for el in body if el.tag.endswith("permEnd") and el.get(qn("w:id")) == "107480931"][0]

    sec2_items = [
        ("2.1 八推进器水下航行器动力学建模",
         "针对紧凑型八推进器水下航行器，开展面向运动控制的动力学建模研究。结合经典物理机理建模与数据驱动（Data-driven）方法，构建涵盖非线性机理模型、降阶解耦模型及数据驱动模型的多种动力学候选模型体系，分析各候选模型的物理表征能力、输入输出形式与结构复杂度；依据推进器在航行器上的空间几何安装位置与矢量朝向，建立单机推力向多自由度广义力与力矩的映射关系，分析控制方向的可达性与推力饱和特性，为后续模型比选与预测控制奠定物理与数学基础。"),
        ("2.2 基于试验数据的水动力参数辨识与预测模型比选",
         "基于水下航行器水池运行实测数据，整理并预处理试验运行数据，提取推进器控制输入与航行器运动响应数据；开展模型辨识与参数拟合，评估模型在不同工况下的泛化能力与预测误差；建立统一的模型评估准则，在独立工况下对多种动力学候选模型的多步递推预测精度、泛化稳定性与单步计算耗时进行横向对比与综合论证，优选出兼顾动力学描述能力与在线计算实时性的预测模型。"),
        ("2.3 面向轨迹跟踪的预测控制与PX4架构接口实现",
         "针对水下精细巡检轨迹跟踪任务，基于所选预测模型设计模型预测控制（MPC）算法；以轨迹跟踪精度与控制平稳性为优化目标，考虑推进器推力饱和等物理约束。围绕航行器轨迹跟踪外环与姿态执行内环的协同问题，深入调研PX4飞控系统的控制分配与接口机制；对比分析“由上层MPC规划运动参考并依托PX4原生内环稳定姿态的分层协同方案”，与“由MPC直接解算多自由度广义控制力矩的集中控制方案”，论证两类接口方案在控制权划分、通信延迟与运行稳定性上的可行性；通过求解耗时统计评估控制方法在规定控制周期内稳定运行的实时性。"),
        ("2.4 基于PX4的闭环仿真与分级实验验证",
         "搭建包含动力学特性、推进器约束和水流扰动的数值仿真环境，开展PX4软件在环（Software-in-the-Loop, SITL）仿真接口联调；设置定点悬停、直线与曲线轨迹及水流扰动等典型工况，以常规控制方法为基准，对比分析轨迹跟踪效果、推力饱和情况与计算求解耗时；重点开展PX4软件在环闭环仿真测试，并视水池试验条件拓展开展典型工况的实物验证。")
    ]

    p_sec2_start_obj = docx.text.paragraph.Paragraph(p_sec2_start, doc)
    set_para_format(p_sec2_start_obj, style_name="Normal", line_val="360", first_line_dxa=0, before_dxa=0, after_dxa=0)
    for r in list(p_sec2_start_obj.runs):
        p_sec2_start_obj._p.remove(r._r)
    r21_h = p_sec2_start_obj.add_run(sec2_items[0][0])
    set_run_font(r21_h, chinese_font="黑体", western_font="Times New Roman", font_size_pt=12.0, bold=False)

    p_curr = p_sec2_start
    p_21_body = create_new_paragraph(doc, sec2_items[0][1], chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False, first_line_dxa=480, line_val="360", style_name="正文格式")
    p_curr.addnext(p_21_body._p)
    p_curr = p_21_body._p

    for h_txt, b_txt in sec2_items[1:]:
        p_h = create_new_paragraph(doc, h_txt, chinese_font="黑体", western_font="Times New Roman", font_size_pt=12.0, bold=False, first_line_dxa=0, before_dxa=0, after_dxa=0, line_val="360", style_name="Normal")
        p_curr.addnext(p_h._p)
        p_curr = p_h._p

        p_b = create_new_paragraph(doc, b_txt, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False, first_line_dxa=480, line_val="360", style_name="正文格式")
        p_curr.addnext(p_b._p)
        p_curr = p_b._p

    elements_to_remove = []
    node = p_curr.getnext()
    while node is not None and node != perm_end_sec2:
        elements_to_remove.append(node)
        node = node.getnext()
    for el in elements_to_remove:
        body.remove(el)

    print("Section 2（主要内容）2.1~2.4 注入完成，小节标题顶格且未加粗，100% 位于 permStart/permEnd 107480931 内部！")

    # ==================== 4. 填充 Section 3：基本要求及成果形式 ====================
    # 更新 Section 3 原生大标题为 "三、毕业论文的基本要求及应完成的成果形式"
    p_h3_node = [el for el in body if el.tag.endswith("p") and any(e.get(qn("w:id")) == "735249041" for e in el.iter(qn("w:permStart")))][0]
    p_h3 = docx.text.paragraph.Paragraph(p_h3_node, doc)
    update_heading_text_safely(p_h3, "三、毕业论文的基本要求及应完成的成果形式")

    p_sec3_start = [el for el in body if any(e.get(qn("w:id")) == "207912164" for e in el.iter(qn("w:permStart")))][0]
    perm_end_sec3 = [el for el in body if el.tag.endswith("permEnd") and el.get(qn("w:id")) == "207912164"][0]

    sec3_texts = [
        "（1）查阅水下航行器动力学、系统辨识、模型预测控制、推力分配及PX4平台适配等方面的中外文献和官方技术资料，归纳研究进展及适用条件，完成文献综述与外文文献翻译。",
        "（2）撰写开题报告，明确研究问题、控制目标、PX4技术路线、实验条件与进度计划，按期参加开题答辩。",
        "（3）建立八推进器运动学与动力学多种候选模型，明确空间推力配置映射关系，形成多模型推导与动力学特性分析文档。",
        "（4）完成试验数据处理与参数拟合，完成多种候选模型的横向对比评估，形成模型对比论证报告、数据处理程序及预测模型实现代码。",
        "（5）完成考虑推进器推力约束的模型预测控制算法设计，打通与PX4系统的控制接口，形成控制程序代码与配置文件。",
        "（6）完成系统仿真与PX4软件在环闭环测试，针对多种典型运动工况输出控制性能对比图表及分析报告；视水池试验条件拓展开展实物测试。",
        "（7）整理模型、程序、配置文件和测试数据，完成毕业论文撰写、修改、查重及毕业答辩。提交符合学校规范的毕业论文、开题报告、文献综述、源程序、PX4配置与接口说明、仿真模型及实验或仿真数据。"
    ]

    p_sec3_start_obj = docx.text.paragraph.Paragraph(p_sec3_start, doc)
    set_para_format(p_sec3_start_obj, style_name="正文格式", line_val="360", first_line_dxa=480)
    for r in list(p_sec3_start_obj.runs):
        p_sec3_start_obj._p.remove(r._r)
    r3_1 = p_sec3_start_obj.add_run(sec3_texts[0])
    set_run_font(r3_1, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False)

    p_curr = p_sec3_start
    for t in sec3_texts[1:]:
        p_new = create_new_paragraph(doc, t, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False, first_line_dxa=480, line_val="360", style_name="正文格式")
        p_curr.addnext(p_new._p)
        p_curr = p_new._p

    elements_to_remove = []
    node = p_curr.getnext()
    while node is not None and node != perm_end_sec3:
        elements_to_remove.append(node)
        node = node.getnext()
    for el in elements_to_remove:
        body.remove(el)

    print("Section 3（基本要求及成果形式）注入完成，100% 位于 permStart/permEnd 207912164 内部！")

    # ==================== 5. 填充 Section 4：进度安排 (Table 2) ====================
    # 更新 Section 4 原生大标题为 "四、毕业论文的进度安排"
    p_h4_node = [el for el in body if el.tag.endswith("p") and any(e.get(qn("w:id")) == "968915720" for e in el.iter(qn("w:permStart")))][0]
    p_h4 = docx.text.paragraph.Paragraph(p_h4_node, doc)
    update_heading_text_safely(p_h4, "四、毕业论文的进度安排")

    p_sec4_start = [el for el in body if any(e.get(qn("w:id")) == "2121286149" for e in el.iter(qn("w:permStart")))][0]
    perm_end_sec4 = [el for el in body if el.tag.endswith("permEnd") and el.get(qn("w:id")) == "2121286149"][0]

    schedule_data = [
        ("序 号", "论文各阶段名称", "时间安排（教学周）"),
        ("1", "课题申报、审核", "2026.7.13 ~ 2026.7.24 （第 20 ~ 21 周）"),
        ("2", "学生选课、选题", "2026.8.4 ~ 2026.8.7 （第 23 周）"),
        ("3", "下达任务书及任务书审核", "2026.8.17 ~ 2026.8.21 （第 25 周）"),
        ("4", "编写开题报告、开题答辩、提交开题报告、审核", "2026.9.16 ~ 2026.10.9 （第 1 ~ 4 周）"),
        ("5", "六自由度动力学建模与空间推力布局分析", "2026.10.12 ~ 2026.11.20 （第 5 ~ 10 周）"),
        ("6", "试验数据处理、动力学建模与预测模型比选", "2026.11.23 ~ 2026.12.25 （第 11 ~ 15 周）"),
        ("7", "中期检查及中期答辩", "2026.12.28 ~ 2027.1.15 （第 16 ~ 18 周）"),
        ("8", "约束MPC与PX4控制接口实现", "2027.2.22 ~ 2027.3.12 （第 1 ~ 3 周）"),
        ("9", "PX4闭环仿真与有条件水池验证", "2027.3.15 ~ 2027.4.2 （第 4 ~ 6 周）"),
        ("10", "毕业论文撰写与图表整理", "2027.4.5 ~ 2027.4.16 （第 7 ~ 8 周）"),
        ("11", "论文查重、导师评阅、专家评阅、学生修改完善论文并定稿、毕业答辩", "2027.4.19 ~ 2027.4.28 （第 9 ~ 10 周）"),
        ("12", "论文评优、论文各类资料（书面版和电子版）整理、归档", "2027.5.3 ~ 2027.5.14 （第 11 ~ 12 周）")
    ]

    tbl_progress = doc.add_table(rows=len(schedule_data), cols=3)
    body.remove(tbl_progress._tbl)
    p_sec4_start.addnext(tbl_progress._tbl)

    tbl_progress.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl_progress._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        r'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    col_widths = [Cm(1.8), Cm(8.6), Cm(5.2)]
    for r_idx, row_vals in enumerate(schedule_data):
        row = tbl_progress.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if r_idx == 0:
            trPr.append(OxmlElement("w:tblHeader"))

        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_c.paragraph_format.line_spacing = 1.15
            p_c.paragraph_format.space_before = Pt(2)
            p_c.paragraph_format.space_after = Pt(2)
            for r in list(p_c.runs):
                p_c._p.remove(r._r)
            run = p_c.add_run(val)
            set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=10.5, bold=False)

    elements_to_remove = []
    node = tbl_progress._tbl.getnext()
    while node is not None and node != perm_end_sec4:
        elements_to_remove.append(node)
        node = node.getnext()
    for el in elements_to_remove:
        body.remove(el)

    print("Section 4（进度安排）Table 2 注入完成，100% 位于 permStart/permEnd 2121286149 内部！")

    # ==================== 6. 填充 Section 5：参考文献 & Section 6：其他要求 ====================
    # 更新 Section 5 原生大标题为 "五、毕业论文应收集的资料及主要参考文献"
    p_h5_node = [el for el in body if el.tag.endswith("p") and any(e.get(qn("w:id")) == "825385557" for e in el.iter(qn("w:permStart")))][0]
    p_h5 = docx.text.paragraph.Paragraph(p_h5_node, doc)
    update_heading_text_safely(p_h5, "五、毕业论文应收集的资料及主要参考文献")

    p_sec5_start = [el for el in body if any(e.get(qn("w:id")) == "1876760563" for e in el.iter(qn("w:permStart")))][0]
    perm_end_sec5 = [el for el in body if el.tag.endswith("permEnd") and el.get(qn("w:id")) == "1876760563"][0]

    intro_text = "应收集的资料包括：PX4水下机型与控制分配文档、推进器安装及标定参数、传感器与试验记录、控制日志以及用于模型和算法验证的仿真数据。"

    references = [
        "[1]  Fossen T I. Handbook of Marine Craft Hydrodynamics and Motion Control[M]. 2nd ed. Chichester, UK: John Wiley & Sons, 2021.",
        "[2]  严浙平, 周佳加. 水下无人航行器控制技术[M]. 北京: 国防工业出版社, 2015.",
        "[3]  Caccia M, Indiveri G, Veruggio G. Modeling and identification of open-frame variable configuration unmanned underwater vehicles[J]. IEEE Journal of Oceanic Engineering, 2000, 25(2): 227-240.",
        "[4]  Heshmati-Alamdari S, Karras G C, Marantos P, et al. A robust model predictive control approach for autonomous underwater vehicles operating in a constrained workspace[C]//2018 IEEE International Conference on Robotics and Automation (ICRA). Brisbane, QLD, Australia: IEEE, 2018: 6183-6188.",
        "[5]  Smallwood D A, Whitcomb L L. Model-based dynamic positioning of underwater robotic vehicles: Theory and experiment[J]. IEEE Journal of Oceanic Engineering, 2004, 29(1): 169-186.",
        "[6]  Johansen T A, Fossen T I. Control allocation: A survey[J]. Automatica, 2013, 49(5): 1087-1103.",
        "[7]  Torrente G, Kaufmann E, Föhn P, et al. Data-driven MPC for quadrotors[J]. IEEE Robotics and Automation Letters, 2021, 6(2): 3769-3776.",
        "[8]  Chu Z, Xiang X, Zhu D, et al. Adaptive trajectory tracking control for remotely operated vehicles considering thruster dynamics and saturation constraints[J]. ISA Transactions, 2020, 100: 28-37.",
        "[9]  孙功武, 苏义鑫, 毛英, 等. 基于模糊逻辑的混合推进ROV多级推力分配策略[J]. 机器人, 2023, 45(4): 472-482.",
        "[10]  Meier L, Honegger D, Pollefeys M. PX4: A node-based multithreaded open source robotics framework for deeply embedded platforms[C]//2015 IEEE International Conference on Robotics and Automation (ICRA). Seattle, WA: IEEE, 2015: 6235-6240.",
        "[11]  PX4 Development Team. Submarines (Unmanned Underwater Vehicles - UUV)[EB/OL]. PX4 Guide. https://docs.px4.io/main/en/frames_sub/ (访问日期: 2026-10-09).",
        "[12]  PX4 Development Team. Control Allocation (Mixing)[EB/OL]. PX4 Guide. https://docs.px4.io/main/en/concept/control_allocation.html (访问日期: 2026-10-09)."
    ]

    p_sec5_start_obj = docx.text.paragraph.Paragraph(p_sec5_start, doc)
    set_para_format(p_sec5_start_obj, style_name="正文格式", line_val="360", first_line_dxa=480)
    for r in list(p_sec5_start_obj.runs):
        p_sec5_start_obj._p.remove(r._r)
    r_intro = p_sec5_start_obj.add_run(intro_text)
    set_run_font(r_intro, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False)

    p_curr = p_sec5_start
    for ref_t in references:
        p_ref = create_ref_paragraph(doc, ref_t, font_size_pt=12.0)
        p_curr.addnext(p_ref._p)
        p_curr = p_ref._p

    # 定位 "六、其他要求" 段落
    p_sec6_heading = None
    node = p_curr.getnext()
    while node is not None and node != perm_end_sec5:
        if node.tag.endswith("p"):
            p_obj = docx.text.paragraph.Paragraph(node, doc)
            if "六、其他要求" in p_obj.text:
                p_sec6_heading = node
                break
        node = node.getnext()

    if p_sec6_heading is not None:
        nodes_between = []
        node = p_curr.getnext()
        while node is not None and node != p_sec6_heading:
            nodes_between.append(node)
            node = node.getnext()
        for el in nodes_between:
            body.remove(el)

    p_sig = None
    node = p_sec6_heading.getnext() if p_sec6_heading is not None else None
    while node is not None and node != perm_end_sec5:
        if node.tag.endswith("p"):
            p_obj = docx.text.paragraph.Paragraph(node, doc)
            if "指导教师签名" in p_obj.text:
                p_sig = node
                break
        node = node.getnext()

    if p_sig is not None:
        p_opt_text = p_sec6_heading.getnext()
        nodes_between_sig = []
        node = p_opt_text.getnext() if p_opt_text is not None else None
        while node is not None and node != p_sig:
            nodes_between_sig.append(node)
            node = node.getnext()
        for el in nodes_between_sig:
            body.remove(el)

    print("Section 5（核准后参考文献 12 篇）与 Section 6 处理完成，100% 位于 permStart/permEnd 1876760563 内部！")

    # ==================== 7. 保存生成文档 ====================
    doc.save(str(target_file_path))
    print(f"文档成功保存至：{target_file_path}")

    # ==================== 8. 执行严格对比与审计 ====================
    run_automated_audit(target_file_path)


def run_automated_audit(built_docx_path):
    print("\n" + "=" * 70)
    print(f">>> 启动自动化验收门禁审计: {Path(built_docx_path).name}")
    print("=" * 70)

    doc_orig = docx.Document(str(TEMPLATE_DOCX))
    doc_built = docx.Document(str(built_docx_path))

    # 1. 权限节点对比（必须 100% 匹配官方模板 24 个 ID）
    orig_starts = [e.get(qn("w:id")) for e in doc_orig._element.iter(qn("w:permStart"))]
    orig_ends = [e.get(qn("w:id")) for e in doc_orig._element.iter(qn("w:permEnd"))]
    built_starts = [e.get(qn("w:id")) for e in doc_built._element.iter(qn("w:permStart"))]
    built_ends = [e.get(qn("w:id")) for e in doc_built._element.iter(qn("w:permEnd"))]

    print(f"模板 permStart 数量: {len(orig_starts)}, 构建后 permStart 数量: {len(built_starts)}")
    print(f"模板 permEnd 数量:   {len(orig_ends)}, 构建后 permEnd 数量:   {len(built_ends)}")
    assert orig_starts == built_starts, f"permStart 不一致! 差异: {set(orig_starts) ^ set(built_starts)}"
    assert orig_ends == built_ends, f"permEnd 不一致! 差异: {set(orig_ends) ^ set(built_ends)}"
    print("[PASS] 官方权限控制节点 (24对 permStart/permEnd) 100% 完整无缺且顺序精确匹配！")

    # 2. 检查大标题文本：必须统一为“毕业论文”
    expected_headings = [
        "一、毕业论文的课题背景",
        "二、毕业论文应完成的主要内容",
        "三、毕业论文的基本要求及应完成的成果形式",
        "四、毕业论文的进度安排",
        "五、毕业论文应收集的资料及主要参考文献",
        "六、其他要求"
    ]
    headings_found = []
    for p in doc_built.paragraphs:
        t = p.text.strip()
        for exp in expected_headings:
            if t == exp:
                headings_found.append(t)
    assert len(headings_found) == 6, f"[FAIL] 大标题不完整或未统一为毕业论文: {headings_found}"
    print(f"[PASS] 全部 6 个大章节标题均已成功统一定格为'毕业论文'！")

    # 3. 严格核查小节标题：绝不加粗 (bold=False)，顶格无缩进
    subheadings_checked = 0
    for p in doc_built.paragraphs:
        t = p.text.strip()
        if any(t.startswith(prefix) for prefix in ["2.1 ", "2.2 ", "2.3 ", "2.4 "]):
            subheadings_checked += 1
            pPr = p._p.find(qn("w:pPr"))
            if pPr is not None:
                ind = pPr.find(qn("w:ind"))
                if ind is not None:
                    assert ind.get(qn("w:firstLine")) is None and ind.get(qn("w:firstLineChars")) is None, f"[FAIL] 小节标题存在缩进: {t}"
            for r in p.runs:
                assert not r.bold, f"[FAIL] 小节标题出现加粗: {t}"
            print(f"[PASS] 小节标题规范校验合格 (顶格无缩进且未加粗): {t[:30]}...")
    assert subheadings_checked == 4, f"[FAIL] 检出小节标题数量异常: {subheadings_checked}"

    # 4. 严格核查正文段落缩进：首行缩进 2 字符 (480 dxa 或 200 chars)
    body_paras_checked = 0
    for p in doc_built.paragraphs:
        t = p.text.strip()
        if t.startswith("随着海洋资源") or t.startswith("然而，水下") or t.startswith("为此，本课题") or \
           t.startswith("针对紧凑型") or t.startswith("基于水下航行器") or t.startswith("针对水下精细") or t.startswith("搭建包含动力学") or \
           t.startswith("（1）") or t.startswith("（2）"):
            body_paras_checked += 1
            pPr = p._p.find(qn("w:pPr"))
            assert pPr is not None, f"[FAIL] 缺少 pPr: {t[:20]}"
            ind = pPr.find(qn("w:ind"))
            assert ind is not None, f"[FAIL] 正文段落未设缩进: {t[:20]}"
            first_line = ind.get(qn("w:firstLine")) or ind.get(qn("w:firstLineChars"))
            assert first_line in ["480", "200"], f"[FAIL] 正文缩进不合规 ({first_line}): {t[:20]}"
            for r in p.runs:
                assert not r.bold, f"[FAIL] 正文段落文字加粗: {t[:20]}"
    assert body_paras_checked >= 9, f"[FAIL] 正文段落校验数量不足: {body_paras_checked}"
    print(f"[PASS] 全部核心正文段落 ({body_paras_checked} 段) 首行缩进 2 字符校验通过！")

    # 5. 严格核查参考文献条目与格式
    ref_count = 0
    for p in doc_built.paragraphs:
        t = p.text.strip()
        if t.startswith("[") and "]" in t[:5]:
            ref_count += 1
            pPr = p._p.find(qn("w:pPr"))
            assert pPr is not None, f"[FAIL] 参考文献缺少 pPr: {t[:20]}"
            ind = pPr.find(qn("w:ind"))
            assert ind is not None, f"[FAIL] 参考文献未设缩进: {t[:20]}"
            assert ind.get(qn("w:hanging")) == "420", f"[FAIL] 参考文献悬挂缩进异常: {t[:20]}"
            for r in p.runs:
                assert not r.bold, f"[FAIL] 参考文献条目被加粗: {t[:20]}"
    assert ref_count == 12, f"[FAIL] 参考文献条目数不是 12 篇 (实际 {ref_count})"
    print(f"[PASS] 参考文献全部 12 条悬挂缩进与未加粗校验 100% 通过！")

    # 6. 禁用词查杀
    full_text = ""
    for p in doc_built.paragraphs:
        full_text += p.text + "\n"
    for t in doc_built.tables:
        for r in t.rows:
            for c in r.cells:
                full_text += c.text + "\n"

    forbidden_terms = ["SwiftROV", "低超调"]
    for term in forbidden_terms:
        count = full_text.count(term)
        assert count == 0, f"[FAIL] 文档中仍残留禁用词 '{term}'，出现次数: {count}！"
        print(f"[PASS] 禁用词 '{term}' 查杀通过 (出现 0 次)！")

    print("=" * 70)
    print(f">>> 验收门禁全部通过！文档 {Path(built_docx_path).name} 格式与规范评级：A+ (完美合规)")
    print("=" * 70)


def main():
    # 生成主任务书、模板名任务书、修改净稿任务书
    build_task_book(OUT_TASK_DOCX_MAIN)
    build_task_book(OUT_TASK_DOCX_TPL)
    build_task_book(OUT_TASK_DOCX_CLEAN)
    print("\n>>> 所有目标任务书文档构建与比对验收圆满完成！")


if __name__ == "__main__":
    main()
