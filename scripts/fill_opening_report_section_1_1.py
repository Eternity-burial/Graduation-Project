# -*- coding: utf-8 -*-
"""
开题报告 Section 1.1 精确填充脚本
功能：
1. 严格以官方空白模板（模板/2毕业设计(论文)开题报告.docx）为基准
2. 填充封面课题名称、学院、专业、姓名、学号及日期
3. 填充 Section 1.1（1．课题来源及研究的目的和意义）三段正文
4. 100% 保持官方全部 25 处可编辑区域 (permStart / permEnd) 完整性与平衡性
5. 严格遵循同济大学排版规范（字体、字号、行距、缩进、上标）
"""

import os
import sys
import re
import shutil
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

TEMPLATE_DOCX = r"模板/2毕业设计(论文)开题报告.docx"
TARGET_DOCX = r"开题报告/2毕业设计(论文)开题报告.docx"
NAMED_DOCX = r"开题报告/开题报告-基于PX4的水下航行器模型控制方法研究.docx"


def set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0, bold=False):
    """精确设置中西文字体、字号与粗体属性"""
    run.font.name = western_font
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), western_font)
    rFonts.set(qn("w:hAnsi"), western_font)
    rFonts.set(qn("w:cs"), western_font)
    rFonts.set(qn("w:eastAsia"), chinese_font)


def set_para_format(p, style_name="正文格式", line_val="360", line_rule="auto", first_line_dxa=480, before_lines=None, after_lines=None, before_pt=None, after_pt=None):
    """精确设置段落样式、行距、首行缩进与段前段后间距"""
    try:
        p.style = style_name
    except Exception:
        pass
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        pPr.append(sp)
    sp.set(qn("w:line"), str(line_val))
    sp.set(qn("w:lineRule"), str(line_rule))
    if before_lines is not None:
        sp.set(qn("w:beforeLines"), str(int(before_lines * 100)))
    if after_lines is not None:
        sp.set(qn("w:afterLines"), str(int(after_lines * 100)))
    if before_pt is not None:
        sp.set(qn("w:before"), str(int(before_pt * 20)))
    if after_pt is not None:
        sp.set(qn("w:after"), str(int(after_pt * 20)))

    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    if first_line_dxa is not None:
        ind.set(qn("w:firstLine"), str(first_line_dxa))
        ind.set(qn("w:firstLineChars"), "0")


def add_text_with_superscripts(paragraph, text, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0):
    """将文本中的参考文献引用 [1], [2], [3] 等精确渲染为小四号上标"""
    parts = re.split(r"(\[\d+\]|\^\{\[\d+\]\})", text)
    for part in parts:
        if not part:
            continue
        m = re.match(r"^(?:\[(\d+)\]|\^\{\[(\d+)\]\})$", part)
        if m:
            ref_num = m.group(1) or m.group(2)
            run = paragraph.add_run(f"[{ref_num}]")
            run.font.superscript = True
            set_run_font(run, chinese_font=chinese_font, western_font=western_font, font_size_pt=font_size_pt)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, chinese_font=chinese_font, western_font=western_font, font_size_pt=font_size_pt)


def set_cell_valign(cell, valign="bottom"):
    """设置表格单元格垂直对齐方式"""
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), valign)


def fill_report():
    print("=" * 70)
    print(">>> 启动开题报告 Section 1.1 精准注入流程")
    print("=" * 70)

    # 1. 备份现有文件
    if os.path.exists(TARGET_DOCX):
        bak_file = TARGET_DOCX + ".bak"
        shutil.copy2(TARGET_DOCX, bak_file)
        print(f"已创建历史版本备份：{bak_file}")

    # 以官方纯净模板为基准加载
    doc = docx.Document(TEMPLATE_DOCX)

    # 2. 保护封面大标题 P5（"毕业设计(论文)开题报告"，36 pt），单倍行距自适应
    p_cover_title = doc.paragraphs[5]
    sp_cover = p_cover_title._p.get_or_add_pPr().find(qn("w:spacing"))
    if sp_cover is None:
        sp_cover = OxmlElement("w:spacing")
        p_cover_title._p.get_or_add_pPr().append(sp_cover)
    sp_cover.set(qn("w:line"), "240")
    sp_cover.set(qn("w:lineRule"), "auto")

    # 3. 填充封面 Table 0（严格保留单元格内的 permStart/permEnd 权限节点）
    t0 = doc.tables[0]
    table_0_entries = [
        (0, 1, "基于PX4的水下航行器模型控制方法研究"),
        (2, 1, "机械工程与机器人学院"),
        (3, 1, "机械设计制造及其自动化"),
        (4, 1, "张卫恒"),
        (4, 3, "2352407"),
    ]
    for r_i, c_i, text_val in table_0_entries:
        cell = t0.rows[r_i].cells[c_i]
        p = cell.paragraphs[0]
        for r in list(p.runs):
            p._p.remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text_val)
        set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=14.0, bold=False)
        set_cell_valign(cell, valign="bottom")

    for row in t0.rows:
        for cell in row.cells:
            set_cell_valign(cell, valign="bottom")

    # 4. 填充封面 Table 1（日期，严格保留列权限节点）
    t1 = doc.tables[1]
    date_entries = [
        (0, "2026"),
        (2, "10"),
        (4, "9"),
    ]
    for c_i, date_val in date_entries:
        cell = t1.rows[0].cells[c_i]
        p = cell.paragraphs[0]
        for r in list(p.runs):
            p._p.remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(date_val)
        set_run_font(run, chinese_font="宋体", western_font="Times New Roman", font_size_pt=14.0)
        set_cell_valign(cell, valign="bottom")

    for cell in t1.rows[0].cells:
        set_cell_valign(cell, valign="bottom")

    # 5. Section 1 骨架定位
    p16 = doc.paragraphs[16] # 一、毕业设计（论文）课题背景
    p17 = doc.paragraphs[17] # 提示文字（包含 permStart 1319923814）
    p18 = doc.paragraphs[18] # 提示文字
    p19 = doc.paragraphs[19] # 1．课题来源及研究的目的和意义
    p20 = doc.paragraphs[20] # 空白占位段落
    p21 = doc.paragraphs[21] # 2．国内外在该方向的研究现状和发展趋势

    # 校准 P16：一、毕业论文课题背景（四号黑体，首行缩进 2 字符，1.5 倍行距）
    p16.runs[2].text = "毕业论文"
    set_para_format(p16, style_name="Normal", line_val="360", line_rule="auto", first_line_dxa=560, before_lines=0.5, after_lines=0.5)
    for r in p16.runs:
        set_run_font(r, chinese_font="黑体", western_font="Times New Roman", font_size_pt=14.0, bold=False)

    # 关键机制：将 permStart 1319923814 从提示段落 P17 迁移绑定至 P19，确保 Section 1 权限 100% 完整
    perm_start_13 = p17._p.find(qn("w:permStart"))
    if perm_start_13 is not None:
        p17._p.remove(perm_start_13)
        pPr_19 = p19._p.get_or_add_pPr()
        pPr_19.addnext(perm_start_13)
        print("已成功将 Section 1 permStart 1319923814 迁移绑定至 P19！")

    # 移除提示段落 P17 与 P18
    p17._element.getparent().remove(p17._element)
    p18._element.getparent().remove(p18._element)

    # 校准 P19（1． 课题来源及研究的目的和意义）：顶格(firstLine=0)，小四号黑体，无加粗，序号与题名空一格
    p19.runs[1].text = "． "
    set_para_format(p19, style_name="条", line_val="360", line_rule="auto", first_line_dxa=0, before_lines=0.5, after_lines=0.5)
    for r in p19.runs:
        set_run_font(r, chinese_font="黑体", western_font="Times New Roman", font_size_pt=12.0, bold=False)

    # 6. 正文草稿内容（经典黄金三段式）
    paragraphs_1_1 = [
        "随着“海洋强国”战略的深入推进以及国家“十四五”规划对高端海洋工程装备的战略部署，海上风电桩基、水下大坝立面及海底能源管线等国家重大水下基础设施的大规模建设与长期服役，对其定期安全巡检与状态评估提出了极为严苛的要求[1]。传统的水下巡检作业主要依赖潜水员人工探查或重型遥控潜水器（ROV）作业，不仅作业窗口期受限、经济成本高昂，且在高危狭窄水域存在显著的人员安全隐患与作业盲区[2]。因此，发展具备高精度自主作业能力的中小型无人水下航行器（UUV），是从根本上摆脱高危人工作业、保障国家重大水下基础设施全生命周期安全运行的核心装备支撑。本课题面向水下基础设施自主巡检的实际工程应用需求，开展水下航行器的自主运动控制方法研究。",
        "在水下大坝立面、风电桩基及管线等典型巡检任务中，受限作业空间对航行器的运动控制性能提出了双重严苛要求：航行器不仅需沿预定空间路径实现厘米级的高精度轨迹跟踪，还必须在近距离对准、壁面观察等工况下满足多自由度姿态独立调节与稳定悬停的需求。常规四推进器或六推进器欠驱动构型中，航行器的位置移动与姿态调节存在强烈的动力学耦合，难以在保持特定俯仰或横滚倾角的同时实现平稳机动；而采用空间矢量布局的紧凑型八推进器航行器具备空间六自由度（6-DOF）全向主动控制能力，为位置跟踪与空间姿态的完全解耦提供了物理基础[3]。然而，该构型在控制实现上面临多重突出瓶颈：其一，流体附加质量、高阶非线性阻尼及重浮力恢复力矩带来复杂的强非线性交叉耦合；其二，受紧凑型结构与电机功率限制，多推进器协同输出时极易引发执行器推力物理饱和（Saturation）与死区效应；其三，常规PID控制难以显式处理多输入多输出（MIMO）系统的执行器硬约束，优化控制方法又面临模型精度与在线计算开销的权衡，且主流开源自主飞控平台（如PX4）对水下航行器的支持仍处于实验性阶段，缺乏面向多矢量推进构型的成熟控制接口与适配机制。",
        "针对上述水下精细巡检任务中多变量耦合、执行器推力饱和及飞控系统适配等关键难题，本课题以八推进器紧凑型水下航行器为研究对象，开展面向控制的动力学建模与参数辨识、考虑物理约束的模型预测控制方法以及PX4开源飞控系统的适配应用研究，旨在实现受限空间下航行器稳定可靠的高精度轨迹跟踪与姿态控制。本研究不仅能为过驱动水下航行器在强非线性与物理执行器饱和约束下的控制系统设计提供理论支持与有效方法参考，同时有助于打通先进约束优化控制在标准化工业级开源飞控上的软硬件工程集成链路，为我国水下基础设施的高精度无人化自主巡检提供具备实际落地价值的技术解决方案。"
    ]

    # 在 P21 之前逐段注入 1.1 正文（小四号宋体，1.5倍行距，首行缩进 2 字符/480 dxa）
    for text in paragraphs_1_1:
        p_body = p21.insert_paragraph_before()
        set_para_format(p_body, style_name="正文格式", line_val="360", line_rule="auto", first_line_dxa=480, before_lines=0, after_lines=0)
        add_text_with_superscripts(p_body, text, chinese_font="宋体", western_font="Times New Roman", font_size_pt=12.0)

    # 移除空占位段落 P20
    p20._element.getparent().remove(p20._element)

    # 7. 保存到命名文件与标准目标文件
    doc.save(NAMED_DOCX)
    print(f"已成功保存至课题专用文档：{NAMED_DOCX}")

    target_saved = False
    try:
        doc.save(TARGET_DOCX)
        print(f"已成功同步保存至标准模板文件：{TARGET_DOCX}")
        target_saved = True
    except PermissionError:
        print(f"[提示] {TARGET_DOCX} 当前正被 WPS/Office 等程序打开占用，已安全写入课题专用文档 {NAMED_DOCX}！")

    # 8. 自动化验收门禁审计
    run_audit(NAMED_DOCX)


def run_audit(doc_path):
    print("\n" + "=" * 70)
    print(">>> 启动权限与排版合规门禁审计")
    print("=" * 70)

    doc_orig = docx.Document(TEMPLATE_DOCX)
    doc_built = docx.Document(doc_path)

    # 1. 检查权限节点是否 100% 完整保留
    orig_starts = [e.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id") for e in doc_orig._element.xpath(".//w:permStart")]
    orig_ends = [e.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id") for e in doc_orig._element.xpath(".//w:permEnd")]
    built_starts = [e.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id") for e in doc_built._element.xpath(".//w:permStart")]
    built_ends = [e.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id") for e in doc_built._element.xpath(".//w:permEnd")]

    missing_s = set(orig_starts) - set(built_starts)
    missing_e = set(orig_ends) - set(built_ends)
    assert len(built_starts) == 25 and len(built_ends) == 25 and not missing_s and not missing_e, f"权限节点缺失！starts={built_starts}, missing_s={missing_s}"
    print("【审计 1 通过】官方全部 25 处可编辑区域 (permStart / permEnd) 100% 完整保留！")

    # 2. 检查权限树闭合顺序
    all_perms = doc_built._element.xpath(".//w:permStart | .//w:permEnd")
    open_stack = {}
    perm_errors = []
    for idx, node in enumerate(all_perms):
        tag = node.tag.split("}")[-1]
        p_id = node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id")
        if tag == "permStart":
            open_stack[p_id] = idx
        elif tag == "permEnd":
            if p_id not in open_stack:
                perm_errors.append(f"permEnd {p_id} 没有对应的 permStart")
            else:
                del open_stack[p_id]
    if open_stack:
        perm_errors.append(f"未闭合的 permStart: {list(open_stack.keys())}")
    assert not perm_errors, f"权限闭合平衡性错误: {perm_errors}"
    print("【审计 2 通过】全部 25 对权限节点在文档树中顺序闭合完全平衡！")

    # 3. 检查 Section 1.1 正文是否 100% 被包裹在可编辑区域 1319923814 内部
    body = doc_built._element.body
    active_perms = set()
    found_section_1_1 = False
    for child in body:
        tag = child.tag.split("}")[-1]
        if tag == "permStart":
            active_perms.add(child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"))
        elif tag == "permEnd":
            active_perms.discard(child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"))
        else:
            for ps in child.xpath(".//w:permStart"):
                active_perms.add(ps.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"))
            for pe in child.xpath(".//w:permEnd"):
                active_perms.discard(pe.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"))

            t = "".join(child.itertext()).strip()
            if "随着“海洋强国”战略" in t:
                assert "1319923814" in active_perms, "第1段未处于可编辑区域 1319923814 内部！"
                found_section_1_1 = True
            if "在水下大坝立面" in t:
                assert "1319923814" in active_perms, "第2段未处于可编辑区域 1319923814 内部！"
            if "针对上述水下精细巡检任务" in t:
                assert "1319923814" in active_perms, "第3段未处于可编辑区域 1319923814 内部！"

    assert found_section_1_1, "未在文档中检测到 Section 1.1 正文！"
    print("【审计 3 通过】Section 1.1 全部三段正文 100% 严密位于可编辑区域 1319923814 内部！")
    print("\n>>> 恭喜！开题报告 1.1 部分已完美填入模板对应可编辑区域！\n")


if __name__ == "__main__":
    fill_report()
