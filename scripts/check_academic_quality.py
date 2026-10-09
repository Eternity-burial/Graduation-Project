# -*- coding: utf-8 -*-
"""
scripts/check_academic_quality.py: 毕业论文写作质量门禁确定性检查脚本

功能与定位：
1. 动态解析 docs/PROJECT_FACTS.md 与 docs/TERMINOLOGY.md 中的 <!-- gate:forbid:关键词 --> 门禁标记；
   两份规则均为必需配置，任一份缺失、不可读、规则标记为空或未解析出有效规则均严格返回 ERROR(2)；
2. 针对指定或默认的学术正文草稿（支持 Markdown 与 Word OOXML DOCX）执行确定性违规扫描；
3. 统一路径授权与安全检查：
   - 授权范围限定为学术正文目录（默认 [开题报告] 目录）；
   - 严格排除系统配置、参考文献、各级归档（归档/archive/历史/history）、图片素材（figures/）及临时文件；
   - 禁止项目根目录、外部绝对路径、越界路径遍历及非授权目录/文件的越权扫描，违规请求明确返回 ERROR(2)；
   - 被排除目录中的文件即使通过完整路径显式指定，亦被严格拦截并返回 ERROR(2)；
4. DOCX 深度扫描：全面覆盖正文段落、表格单元格、嵌套表格、节页眉与节页脚（含默认/首页/奇偶页）；
5. 严格的三态退出语义（PASS 0 / FAIL 1 / ERROR 2）：
   - 杜绝虚假 PASS：未发现正文、正文为空内容、文件损坏、规则缺失均返回 ERROR(2)；
   - 确定性违规返回 FAIL(1)；全部合规且实际完成扫描返回 PASS(0)。

覆盖限制说明 (Coverage Limitations)：
本检查器覆盖 DOCX 的正文段落、各级表格与嵌套表格、节页眉页脚（含首页与奇偶页）。
对于 Word OOXML 的独立文本框 (w:txbxContent)、批注与修订 (w:comments)、SmartArt
及嵌入对象中的非流式特殊结构，因其非常规正文排版元素，不在本自动化门禁的确定性读取范围内，
需在最终交付前由人工审阅核验。
"""

import sys
import re
import argparse
from pathlib import Path

# 确保控制台 UTF-8 输出
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# 必需的规则文件清单（缺一不可）
REQUIRED_RULE_FILES = ["PROJECT_FACTS.md", "TERMINOLOGY.md"]

# 授权的学术正文目录名称（仅允许扫描这些目录下的正文）
AUTHORIZED_DOC_DIR_NAMES = {"开题报告"}

# 绝对排除扫描的目录名称与特征（黑名单，涵盖根目录与各级子目录）
IGNORED_DIR_NAMES = {
    ".git",
    ".github",
    ".agents",
    "docs",
    "参考",
    "资料",
    "模板",
    "PX4_INDI_Research",
    "选题",
    "归档",
    "archive",
    "历史",
    "history",
    "figures",
    "修改建议",
    "__pycache__",
}

# 绝对排除扫描的文件名称
IGNORED_FILE_NAMES = {
    "AGENTS.md",
    "PROJECT_NOW.md",
    "PROJECT_STATUS.md",
    "references.bib",
    "references.ris",
}

# 允许的学术正文扩展名
ALLOWED_EXTENSIONS = {".md", ".docx", ".txt"}


class QualityCheckResult:
    """质量门禁执行结果对象，封装三态结果并支持元组解包 (passed, violations, rules)"""

    def __init__(self, status: str, violations: list, rules: list, scanned_files: list, errors: list):
        self.status = status  # "PASS", "FAIL", "ERROR"
        self.violations = violations
        self.rules = rules
        self.scanned_files = scanned_files
        self.errors = errors

    @property
    def passed(self) -> bool:
        return self.status == "PASS"

    @property
    def exit_code(self) -> int:
        if self.status == "PASS":
            return 0
        elif self.status == "FAIL":
            return 1
        else:
            return 2

    def __iter__(self):
        """兼容原有 (passed, violations, rules) 解包语法"""
        return iter((self.passed, self.violations, self.rules))


def load_gate_rules(project_root: Path):
    """
    从 docs/PROJECT_FACTS.md 与 docs/TERMINOLOGY.md 提取所有 <!-- gate:forbid:KEYWORD --> 规则。
    两份文件均为必需配置。任意一份缺失、不可读取或未包含有效门禁规则，均记录致命错误。
    返回: (rules, errors)
    """
    rules = []
    errors = []
    seen_keywords = set()
    docs_dir = project_root / "docs"

    if not docs_dir.exists() or not docs_dir.is_dir():
        errors.append(f"必需规则目录不存在: {docs_dir}")
        return rules, errors

    for doc_name in REQUIRED_RULE_FILES:
        doc_path = docs_dir / doc_name
        if not doc_path.exists() or not doc_path.is_file():
            errors.append(f"必需门禁规则文件缺失: docs/{doc_name}")
            continue

        try:
            content = doc_path.read_text(encoding="utf-8")
        except Exception as e:
            errors.append(f"无法读取门禁规则文件 docs/{doc_name}: {e}")
            continue

        lines = content.splitlines()
        current_id = "UNKNOWN"
        file_rule_count = 0

        for idx, line in enumerate(lines, start=1):
            stripped_line = line.strip()

            # 匹配形如 ### F-001： 或 ### T-001： 的编号（修复原正则 [F|T] 误接受竖线的缺陷）
            id_match = re.match(r"^###\s+([FT]-\d+)", stripped_line)
            if id_match:
                current_id = id_match.group(1)

            # 匹配 <!-- gate:forbid:关键词 -->
            gate_match = re.search(r"<!--\s*gate:forbid:(.*?)\s*-->", line)
            if gate_match:
                keyword = gate_match.group(1).strip()
                if not keyword:
                    errors.append(f"门禁规则文件 docs/{doc_name} 第 {idx} 行禁词标记为空！")
                    continue

                if current_id == "UNKNOWN":
                    errors.append(f"门禁规则文件 docs/{doc_name} 第 {idx} 行禁词 \"{keyword}\" 缺少有效规则编号（未处于合法 ### F-xxx 或 ### T-xxx 标题下）！")
                    continue

                # 重复规则处理：若同规则源内重复出现同一关键词，进行去重并记录
                rule_key = (doc_name, keyword)
                if rule_key in seen_keywords:
                    continue
                seen_keywords.add(rule_key)

                rules.append({
                    "keyword": keyword,
                    "rule_id": current_id,
                    "source_file": doc_name,
                    "rule_line": idx,
                })
                file_rule_count += 1

        if file_rule_count == 0:
            errors.append(f"门禁规则文件 docs/{doc_name} 未包含任何有效 <!-- gate:forbid:... --> 规则")

    return rules, errors


def is_blacklisted(file_path: Path, project_root: Path) -> bool:
    """检查文件或目录是否位于严格忽略的黑名单路径中（覆盖任意层级目录）"""
    try:
        rel = file_path.resolve().relative_to(project_root.resolve())
    except ValueError:
        # 不在 project_root 内部，视为越界
        return True

    parts = rel.parts
    if not parts:
        return True

    # 检查相对路径中所有目录层级：若命中黑名单目录或包含归档/历史字样，直接过滤
    dir_parts = parts[:-1] if file_path.is_file() else parts
    for part in dir_parts:
        if part in IGNORED_DIR_NAMES or "归档" in part or "历史" in part or "archive" in part.lower() or "history" in part.lower():
            return True
        if part.startswith(("_test_sandbox_", "_temp_")):
            return True

    # 检查具体文件
    if file_path.is_file():
        if file_path.name in IGNORED_FILE_NAMES:
            return True
        if file_path.suffix.lower() in [".bak", ".tmp", ".log"]:
            return True
        if file_path.name.startswith("~$"):
            return True
        if file_path.name.startswith("基于PX4"):
            # 历史任务书参考底稿
            return True

    return False


def validate_target_path(target_path: Path, project_root: Path) -> tuple[bool, str]:
    """
    统一路径授权与合法性检查。
    验证路径存在性、边界授权、黑名单排除及文件类型。
    返回: (is_valid, error_message)
    """
    resolved = target_path.resolve()
    if not resolved.exists():
        return False, f"指定的目标路径不存在: {target_path}"

    try:
        rel = resolved.relative_to(project_root.resolve())
    except ValueError:
        return False, f"目标路径超出项目授权范围: {target_path}"

    if resolved == project_root.resolve():
        return False, f"禁止将整个项目根目录作为受检目标，请指定具体正文文件或 [开题报告] 目录: {target_path}"

    if len(rel.parts) == 0 or rel.parts[0] not in AUTHORIZED_DOC_DIR_NAMES:
        return False, f"指定路径超出学术正文受检授权范围（仅允许检查 {list(AUTHORIZED_DOC_DIR_NAMES)} 目录及其正文文件）: {target_path}"

    if is_blacklisted(resolved, project_root):
        return False, f"目标路径属于受保护黑名单或非学术正文范围: {target_path}"

    if resolved.is_file():
        if resolved.suffix.lower() not in ALLOWED_EXTENSIONS:
            return False, f"不支持的文档类型: {resolved.suffix} (仅支持 .md, .txt, .docx): {target_path}"

    return True, ""


def _extract_table_lines(tbl, prefix="表格"):
    """递归提取表格及其单元格内嵌套表格的文本行，自动对合并单元格去重"""
    lines = []
    seen_cells = set()
    for r_idx, row in enumerate(tbl.rows, start=1):
        for c_idx, cell in enumerate(row.cells, start=1):
            cell_id = cell._tc
            if cell_id in seen_cells:
                continue
            seen_cells.add(cell_id)

            cell_loc = f"{prefix} (行{r_idx}, 列{c_idx})"
            # 提取单元格内的段落文本
            for p_idx, p in enumerate(cell.paragraphs, start=1):
                t = p.text.strip()
                if t:
                    loc = f"{cell_loc} 段落{p_idx}" if len(cell.paragraphs) > 1 else cell_loc
                    lines.append((None, t, loc))

            # 递归提取嵌套表格
            for sub_t_idx, sub_tbl in enumerate(cell.tables, start=1):
                sub_lines = _extract_table_lines(
                    sub_tbl,
                    prefix=f"{cell_loc} -> 嵌套表格{sub_t_idx}"
                )
                lines.extend(sub_lines)
    return lines


def _scan_header_or_footer(hf, desc, lines):
    """扫描页眉或页脚中的段落与表格"""
    if hf is None:
        return
    for p_idx, p in enumerate(hf.paragraphs, start=1):
        t = p.text.strip()
        if t:
            lines.append((None, t, f"{desc} (段落{p_idx})"))
    for t_idx, tbl in enumerate(hf.tables, start=1):
        lines.extend(_extract_table_lines(tbl, prefix=f"{desc} (表格{t_idx})"))


def extract_file_lines(file_path: Path):
    """
    从目标文件中提取文本行列表。
    - Markdown / Text: 按 UTF-8 解码提取行
    - DOCX: 深度覆盖节页眉、节页脚（含首页/奇偶页）、正文段落及各级表格与嵌套表格
    返回: list of tuple (line_num or None, text, location_str)
    """
    ext = file_path.suffix.lower()
    if ext in [".md", ".txt"]:
        content = file_path.read_text(encoding="utf-8")
        lines = []
        for idx, line in enumerate(content.splitlines(), start=1):
            stripped = line.strip()
            if stripped:
                lines.append((idx, stripped, f"第 {idx} 行"))
        return lines

    elif ext == ".docx":
        import docx
        doc = docx.Document(str(file_path))
        lines = []

        # 1. 扫描所有节的页眉与页脚（全面覆盖默认、首页及奇偶页）
        for s_idx, sec in enumerate(doc.sections, start=1):
            _scan_header_or_footer(sec.header, f"页眉 (第{s_idx}节)", lines)
            _scan_header_or_footer(sec.footer, f"页脚 (第{s_idx}节)", lines)
            if getattr(sec, "different_first_page_header_footer", False):
                _scan_header_or_footer(getattr(sec, "first_page_header", None), f"首页页眉 (第{s_idx}节)", lines)
                _scan_header_or_footer(getattr(sec, "first_page_footer", None), f"首页页脚 (第{s_idx}节)", lines)
            if getattr(doc.settings, "odd_and_even_pages_header_footer", False):
                _scan_header_or_footer(getattr(sec, "even_page_header", None), f"偶数页页眉 (第{s_idx}节)", lines)
                _scan_header_or_footer(getattr(sec, "even_page_footer", None), f"偶数页页脚 (第{s_idx}节)", lines)

        # 2. 扫描正文段落
        for p_idx, p in enumerate(doc.paragraphs, start=1):
            t = p.text.strip()
            if t:
                lines.append((p_idx, t, f"段落 {p_idx}"))

        # 3. 扫描正文表格（包含嵌套表格）
        for t_idx, tbl in enumerate(doc.tables, start=1):
            lines.extend(_extract_table_lines(tbl, prefix=f"表格 {t_idx}"))

        return lines

    else:
        raise ValueError(f"不支持的文档类型: {ext} (仅支持 .md, .txt, .docx)")


def is_academic_draft_name(name: str) -> bool:
    """判断文件名称是否符合学术正文草稿特征（用于默认发现排除笔记/梳理材料）"""
    stem = Path(name).stem
    # 明确排除非正文梳理、笔记或建议
    exclude_keywords = ["梳理", "笔记", "大纲", "建议", "修改", "对比", "分析"]
    if any(k in stem for k in exclude_keywords):
        return False
    # 匹配正文草稿命名特征
    include_keywords = ["开题报告", "正文", "起草", "草稿", "chapter", "thesis"]
    return any(k in stem.lower() for k in include_keywords)


def find_default_targets(project_root: Path):
    """
    递归查找默认需要受检的学术正文草稿及交付文件。
    授权范围：[开题报告] 目录及其所有正文子目录
    排除：各级归档、素材、历史任务书底稿、脚本、参考文献及非正文辅助分析文件
    """
    targets = []
    report_dir = project_root / "开题报告"
    if not report_dir.exists() or not report_dir.is_dir():
        return targets

    for p in sorted(report_dir.rglob("*")):
        if not p.is_file():
            continue
        if is_blacklisted(p, project_root):
            continue
        ext = p.suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            continue
        if is_academic_draft_name(p.name):
            targets.append(p)

    return sorted(targets)


def run_quality_check(target_paths=None, project_root=None) -> QualityCheckResult:
    """
    程序化执行质量门禁检查，返回确定性 QualityCheckResult 对象
    """
    if project_root is None:
        project_root = Path(__file__).resolve().parent.parent

    errors = []
    rules, rule_errors = load_gate_rules(project_root)
    errors.extend(rule_errors)

    targets = []
    if target_paths is not None:
        if len(target_paths) == 0:
            errors.append("指定了 --target 参数但未提供任何有效目标路径（若需使用默认范围请不要指定 --target 参数）！")
        else:
            for p in target_paths:
                path_obj = Path(p)
                valid, msg = validate_target_path(path_obj, project_root)
                if not valid:
                    errors.append(msg)
                    continue

                resolved_path = path_obj.resolve()
                if resolved_path.is_file():
                    targets.append(resolved_path)
                elif resolved_path.is_dir():
                    found = []
                    for f in sorted(resolved_path.rglob("*")):
                        if not f.is_file():
                            continue
                        if is_blacklisted(f, project_root):
                            continue
                        if f.suffix.lower() in ALLOWED_EXTENSIONS:
                            found.append(f)
                    if not found:
                        errors.append(f"指定目录下未找到任何有效的正文文件: {p}")
                    else:
                        targets.extend(found)
    else:
        targets = find_default_targets(project_root)
        if not targets:
            errors.append("在授权目录 [开题报告] 中未找到任何默认的正文草稿或交付文件！")

    # 去重并排序
    targets = sorted(list(set(targets)))

    if not targets and not errors:
        errors.append("受检正文文件集合为空（未实际扫描任何正文文件）！")

    all_violations = []
    scanned_files = []

    for target in targets:
        try:
            lines = extract_file_lines(target)
        except Exception as e:
            errors.append(f"无法读取或解析文件 {target}: {e}")
            continue

        if not lines:
            errors.append(f"文件存在但无可检查的有效正文内容（空文件或无文本段落）: {target}")
            continue

        scanned_files.append(target)

        for line_num, line_text, loc_desc in lines:
            for rule in rules:
                kw = rule["keyword"]
                if kw in line_text:
                    all_violations.append({
                        "file": target,
                        "location": loc_desc,
                        "line_num": line_num,
                        "content": line_text,
                        "keyword": kw,
                        "rule_id": rule["rule_id"],
                        "source": rule["source_file"],
                    })

    # 判定三态结果
    if errors:
        status = "ERROR"
    elif all_violations:
        status = "FAIL"
    elif scanned_files:
        status = "PASS"
    else:
        status = "ERROR"
        errors.append("实际扫描文件数量为 0！")

    return QualityCheckResult(status, all_violations, rules, scanned_files, errors)


def main():
    parser = argparse.ArgumentParser(description="学术写作质量门禁确定性检查工具")
    parser.add_argument("--target", "-t", nargs="*", default=None, help="指定受检目标文件或目录（默认扫描开题报告正文草稿）")
    parser.add_argument("--root", "-r", default=None, help="指定项目根目录（可选）")
    args = parser.parse_args()

    project_root = Path(args.root).resolve() if args.root else Path(__file__).resolve().parent.parent

    # 显式传递 --target 但未提供实际参数时，明确报错拦截
    if args.target is not None and len(args.target) == 0:
        print(f"==================================================")
        print(f" 学术写作流程确定性质量门禁 (Quality Gate Check)")
        print(f" 根目录: {project_root}")
        print(f"==================================================")
        print("\n[ERROR] 质量门禁执行失败！命令行参数 --target 未指定具体目标路径（若需使用默认范围请不要传递 --target 参数）\n")
        sys.exit(2)

    print(f"==================================================")
    print(f" 学术写作流程确定性质量门禁 (Quality Gate Check)")
    print(f" 根目录: {project_root}")
    print(f"==================================================")

    result = run_quality_check(args.target, project_root)

    print(f"[INFO] 已加载 {len(result.rules)} 条生效门禁规则:")
    for r in result.rules:
        print(f"  - [{r['rule_id']}] 严禁词: \"{r['keyword']}\" (来自 {r['source_file']})")
    print("--------------------------------------------------")

    if result.scanned_files:
        print(f"[INFO] 本次实际扫描正文文件 (共 {len(result.scanned_files)} 个):")
        for idx, f in enumerate(result.scanned_files, start=1):
            try:
                rel = f.relative_to(project_root)
            except ValueError:
                rel = f
            print(f"  {idx}. {rel}")
        print("--------------------------------------------------")

    if result.status == "PASS":
        print("\n[PASS] 质量门禁检查通过！")
        print(f"本次实际扫描 {len(result.scanned_files)} 个正文文件，未发现任何阻断级违规。\n")
        sys.exit(0)
    elif result.status == "FAIL":
        print(f"\n[FAIL] 质量门禁检查未通过！发现 {len(result.violations)} 处阻断级违规:\n")
        for idx, v in enumerate(result.violations, start=1):
            try:
                rel = v["file"].relative_to(project_root)
            except ValueError:
                rel = v["file"]
            print(f"  {idx}. [{v['rule_id']}] 命中禁词: \"{v['keyword']}\"")
            print(f"     文件: {rel}")
            print(f"     位置: {v['location']}")
            print(f"     正文摘录: \"{v['content']}\"")
            print(f"     规则来源: docs/{v['source']}")
            print()
        print("【整改要求】：请立即根据 docs/PROJECT_FACTS.md 与 docs/TERMINOLOGY.md 修正上述用词后再行提交验收！\n")
        sys.exit(1)
    else:  # ERROR
        print(f"\n[ERROR] 质量门禁执行失败！发现 {len(result.errors)} 处配置或系统致命异常:\n")
        for idx, err in enumerate(result.errors, start=1):
            print(f"  {idx}. {err}")
        print()
        sys.exit(2)


if __name__ == "__main__":
    main()
